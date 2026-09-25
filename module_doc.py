import os
import shutil
import asyncio
from datetime import datetime
import flet as ft
import httpx
from PIL import Image, ImageOps

import config
import core


class ModulDocMixin:
    def _inicjalizuj_modul_dok(self):
        # Bezpieczne granice kadrowania dopasowane do ekranu telefonu
        self.MAX_SZEROKOSC_ROBOCZA_DOC = 330.0
        self.MAX_WYSOKOSC_ROBOCZA_DOC = 460.0
        self.UCHWYT_ROZMIAR_DOK_DOC = 40

        self.SZEROKOSC_DOK_DOC = int(self.MAX_SZEROKOSC_ROBOCZA_DOC)
        self.WYSOKOSC_DOK_DOC = int(self.MAX_WYSOKOSC_ROBOCZA_DOC)

        self.crop_doc_x1 = 0.0
        self.crop_doc_y1 = 0.0
        self.crop_doc_x2 = float(self.SZEROKOSC_DOK_DOC)
        self.crop_doc_y2 = float(self.WYSOKOSC_DOK_DOC)
        self.kadr_doc_zmieniony = False

        self.historia_dok = []
        self.aktywny_filtr_dok = None
        self.zdjecie_dok = {"sciezka": None}
        self.ostatni_wynik_dok = {"tekst": ""}

    def _inicjalizuj_modul_dokument(self):
        self._inicjalizuj_modul_dok()

        # 1. Podgląd na ekranie głównym modułu Dokument
        self.podglad_dok = ft.Image(
            src=config.PUSTY_OBRAZ,
            fit="contain",
            width=320,
            height=240
        )

        # 2. Obraz na pełnym ekranie kadrowania
        self.img_pelny_podglad_dok = ft.Image(
            src=config.PUSTY_OBRAZ,
            fit="fill",
            width=self.SZEROKOSC_DOK_DOC,
            height=self.WYSOKOSC_DOK_DOC
        )

        self.maska_doc_gora = ft.Container(bgcolor=ft.Colors.BLACK54, top=0, left=0, width=self.SZEROKOSC_DOK_DOC, height=0)
        self.maska_doc_dol = ft.Container(bgcolor=ft.Colors.BLACK54, bottom=0, left=0, width=self.SZEROKOSC_DOK_DOC, height=0)
        self.maska_doc_lewo = ft.Container(bgcolor=ft.Colors.BLACK54, top=0, bottom=0, left=0, width=0)
        self.maska_doc_prawo = ft.Container(bgcolor=ft.Colors.BLACK54, top=0, bottom=0, right=0, width=0)

        self.strefa_srodka_dok = ft.GestureDetector(
            content=ft.Container(
                border=ft.Border.all(2.0, ft.Colors.BLUE_ACCENT),
                bgcolor=ft.Colors.with_opacity(0.08, ft.Colors.BLUE_ACCENT)
            ),
            drag_interval=10,
            on_pan_update=lambda e: self._przesun_caly_kadr_dok(*self._pobierz_deltas_dok(e)),
            top=0, left=0, width=self.SZEROKOSC_DOK_DOC, height=self.WYSOKOSC_DOK_DOC
        )

        def stworz_uchwyt_dok():
            return ft.Container(
                alignment=ft.Alignment(0, 0),
                content=ft.Container(
                    width=26, height=26,
                    bgcolor=ft.Colors.BLUE_ACCENT,
                    border_radius=13,
                    border=ft.Border.all(2.0, ft.Colors.WHITE)
                ),
                width=self.UCHWYT_ROZMIAR_DOK_DOC,
                height=self.UCHWYT_ROZMIAR_DOK_DOC
            )

        self.uchwyt_doc_lt = ft.GestureDetector(content=stworz_uchwyt_dok(), drag_interval=10, on_pan_update=lambda e: self._przesun_uchwyt_dok("lt", *self._pobierz_deltas_dok(e)))
        self.uchwyt_doc_rt = ft.GestureDetector(content=stworz_uchwyt_dok(), drag_interval=10, on_pan_update=lambda e: self._przesun_uchwyt_dok("rt", *self._pobierz_deltas_dok(e)))
        self.uchwyt_doc_lb = ft.GestureDetector(content=stworz_uchwyt_dok(), drag_interval=10, on_pan_update=lambda e: self._przesun_uchwyt_dok("lb", *self._pobierz_deltas_dok(e)))
        self.uchwyt_doc_rb = ft.GestureDetector(content=stworz_uchwyt_dok(), drag_interval=10, on_pan_update=lambda e: self._przesun_uchwyt_dok("rb", *self._pobierz_deltas_dok(e)))

        self.ramka_kadrowania_dok = ft.Container(
            content=ft.Stack([
                self.img_pelny_podglad_dok,
                self.maska_doc_gora, self.maska_doc_dol, self.maska_doc_lewo, self.maska_doc_prawo,
                self.strefa_srodka_dok,
                self.uchwyt_doc_lt, self.uchwyt_doc_rt, self.uchwyt_doc_lb, self.uchwyt_doc_rb
            ]),
            width=self.SZEROKOSC_DOK_DOC,
            height=self.WYSOKOSC_DOK_DOC,
            alignment=ft.Alignment(0, 0)
        )

        # Górna belka
        self.btn_obrot_l_dok = ft.IconButton(
            icon=ft.Icons.ROTATE_LEFT, icon_color=ft.Colors.WHITE, icon_size=28,
            width=50, height=50,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, shape=ft.RoundedRectangleBorder(radius=10)),
            tooltip="Obróć w lewo (90°)",
            on_click=lambda e: asyncio.create_task(self.obroc_dok(90))
        )
        self.btn_doc_f_bw = ft.Button(
            content=ft.Text("B&W", size=12, weight=ft.FontWeight.BOLD),
            height=50, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._przelacz_filtr_pelny_dok("bw"))
        )
        self.btn_doc_f_szary = ft.Button(
            content=ft.Text("GRY", size=12, weight=ft.FontWeight.BOLD),
            height=50, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._przelacz_filtr_pelny_dok("szary"))
        )
        self.btn_doc_f_wyostrz = ft.Button(
            content=ft.Text("SHP", size=12, weight=ft.FontWeight.BOLD),
            height=50, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._przelacz_filtr_pelny_dok("wyostrz"))
        )
        self.btn_obrot_r_dok = ft.IconButton(
            icon=ft.Icons.ROTATE_RIGHT, icon_color=ft.Colors.WHITE, icon_size=28,
            width=50, height=50,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, shape=ft.RoundedRectangleBorder(radius=10)),
            tooltip="Obróć w prawo (90°)",
            on_click=lambda e: asyncio.create_task(self.obroc_dok(-90))
        )

        self.wiersz_filtrow_dok = ft.Row([
            self.btn_obrot_l_dok,
            self.btn_doc_f_bw,
            self.btn_doc_f_szary,
            self.btn_doc_f_wyostrz,
            self.btn_obrot_r_dok
        ], spacing=4)

        # Dolne przyciski
        self.btn_doc_anuluj = ft.Button(
            content=ft.Text("Anuluj", size=13, weight=ft.FontWeight.BOLD),
            height=52, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_900, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._zamknij_pelny_ekran_dok(zapisz=False))
        )
        self.btn_doc_cofnij = ft.Button(
            content=ft.Text("Cofnij", size=13, weight=ft.FontWeight.BOLD),
            height=52, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREY_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            disabled=True,
            on_click=lambda e: self._cofnij_krok_pelny_dok()
        )
        self.btn_doc_zatwierdz = ft.Button(
            content=ft.Text("Zatwierdź", size=13, weight=ft.FontWeight.BOLD),
            height=52, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.BLUE_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._zatwierdz_krok_pelny_dok())
        )
        self.btn_doc_wyslij = ft.Button(
            content=ft.Text("Wyślij", size=13, weight=ft.FontWeight.BOLD),
            height=52, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREEN_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=10), padding=0),
            on_click=lambda e: asyncio.create_task(self._wyslij_pelny_ekran_dok())
        )
        self.wiersz_akcji_pelnych_dok = ft.Row([
            self.btn_doc_anuluj,
            self.btn_doc_cofnij,
            self.btn_doc_zatwierdz,
            self.btn_doc_wyslij
        ], spacing=4)

        # Dedykowany pełny kontener kadrowania
        self.widok_kadrowania_dok = ft.Container(
            content=ft.Column([
                self.wiersz_filtrow_dok,
                ft.Container(
                    content=self.ramka_kadrowania_dok,
                    alignment=ft.Alignment(0, 0),
                    expand=True,
                    padding=6
                ),
                self.wiersz_akcji_pelnych_dok
            ], spacing=6, alignment=ft.MainAxisAlignment.SPACE_BETWEEN, horizontal_alignment=ft.CrossAxisAlignment.CENTER),
            padding=6,
            expand=True,
            visible=False
        )

        # Kontrolki ekranu głównego modułu Dokument
        self.status_dok = ft.Text("Wybierz zdjęcie dokumentu biurowego lub zrób nowe.", size=13, color=ft.Colors.GREY_300, text_align=ft.TextAlign.CENTER)
        self.pasek_dok = ft.ProgressBar(visible=False, color=ft.Colors.LIGHT_BLUE_ACCENT)

        self.btn_foto_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.PHOTO_LIBRARY), ft.Text("ZDJĘCIA")], alignment=ft.MainAxisAlignment.CENTER),
            height=55, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.GREEN_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            on_click=lambda e: asyncio.create_task(self.otworz_galerie_dok())
        )
        self.btn_aparat_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.CAMERA_ALT), ft.Text("Zrób zdjęcie")], alignment=ft.MainAxisAlignment.CENTER),
            height=55, expand=True,
            style=ft.ButtonStyle(bgcolor=ft.Colors.BLUE_900, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            on_click=lambda e: asyncio.create_task(self.ui.otworz_aparat_dla("dokument"))
        )
        self.wiersz_foto_dok = ft.Row([self.btn_foto_dok, self.btn_aparat_dok], spacing=10)

        self.kontener_podgladu_dok = ft.Container(
            content=ft.GestureDetector(
                content=self.podglad_dok,
                on_tap=lambda e: asyncio.create_task(self.otworz_pelny_podglad_dok(e))
            ),
            alignment=ft.Alignment(0, 0),
            height=250,
            border=ft.Border.all(1, ft.Colors.GREY_800),
            border_radius=8,
            visible=False
        )

        self.btn_otworz_kadr_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.CROP, size=20), ft.Text("Dopasuj kadr / Filtry", weight=ft.FontWeight.BOLD)], alignment=ft.MainAxisAlignment.CENTER),
            height=48, visible=False,
            style=ft.ButtonStyle(bgcolor=ft.Colors.BLUE_900, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            on_click=lambda e: asyncio.create_task(self.otworz_pelny_podglad_dok(e))
        )
        self.btn_start_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.DOCUMENT_SCANNER), ft.Text("Odczytaj dokument (1:1)")], alignment=ft.MainAxisAlignment.CENTER),
            height=48, visible=False,
            style=ft.ButtonStyle(bgcolor=ft.Colors.LIGHT_BLUE_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            on_click=lambda e: asyncio.create_task(self.analizuj_dokument_dok(e))
        )
        self.btn_usun_foto_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.DELETE_OUTLINE, size=18), ft.Text("Usuń wybrane zdjęcie", size=12)], alignment=ft.MainAxisAlignment.CENTER),
            style=ft.ButtonStyle(bgcolor=ft.Colors.RED_900, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            visible=False,
            on_click=lambda e: self.usun_wybrane_zdjecie_dok()
        )
        self.btn_wroc_wynik_dok = ft.Button(
            content=ft.Row([ft.Icon(ft.Icons.VISIBILITY), ft.Text("Pokaż ostatni odczytany tekst")], alignment=ft.MainAxisAlignment.CENTER),
            height=48, visible=False,
            style=ft.ButtonStyle(bgcolor=ft.Colors.TEAL_800, color=ft.Colors.WHITE, shape=ft.RoundedRectangleBorder(radius=8)),
            on_click=lambda e: self.ui.bezpiecznie_otworz_dialog(self.dlg_wynik_dok)
        )

        self.btn_konsola_dok = ft.IconButton(icon=ft.Icons.TERMINAL, tooltip="Konsola zdarzeń (logi)", on_click=self.ui.otworz_konsole)

        pasek_tytulu_dok = ft.Row([
            ft.Row([
                ft.IconButton(ft.Icons.ARROW_BACK, tooltip="Menu Główne", on_click=lambda e: asyncio.create_task(self.przelacz_widok("menu"))),
                ft.Column([
                    ft.Text("ocrLmm Mobile", size=18, weight=ft.FontWeight.BOLD, color=ft.Colors.LIGHT_BLUE_400),
                    ft.Text("Skaner Dokumentów (Word/Excel/TXT)", size=11, color=ft.Colors.GREY_400)
                ], spacing=1)
            ]),
            self.btn_konsola_dok
        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN)

        self.kolumna_glowna_dok = ft.Column([
            pasek_tytulu_dok,
            ft.Divider(height=10, color=ft.Colors.TRANSPARENT),
            self.wiersz_foto_dok,
            self.pasek_dok,
            self.status_dok,
            self.kontener_podgladu_dok,
            self.btn_otworz_kadr_dok,
            self.btn_start_dok,
            self.btn_wroc_wynik_dok,
            self.btn_usun_foto_dok
        ], horizontal_alignment=ft.CrossAxisAlignment.STRETCH, spacing=10)

        self.widok_dok = ft.Column([
            self.kolumna_glowna_dok,
            self.widok_kadrowania_dok
        ], horizontal_alignment=ft.CrossAxisAlignment.STRETCH, spacing=0, visible=False)

        self.widok_dokument = self.widok_dok

        # Dialog edytora tekstu z eksportem (dopasowany do szerokości ekranu)
        self.txt_edytor_dok = ft.TextField(
            multiline=True,
            min_lines=14,
            max_lines=22,
            text_size=12,
            expand=True
        )
        self.dlg_wynik_dok = ft.AlertDialog(
            modal=True,
            title=ft.Row([
                ft.Text("Odczytany dokument", size=16, weight=ft.FontWeight.BOLD),
                ft.IconButton(ft.Icons.CLOSE, on_click=lambda e: self.page.pop_dialog())
            ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
            content=ft.Container(
                content=self.txt_edytor_dok,
                height=420,
                expand=True
            ),
            actions=[
                ft.Row([
                    ft.Button(
                        content=ft.Text("Kopiuj", size=12),
                        expand=True,
                        style=ft.ButtonStyle(padding=0),
                        on_click=lambda e: asyncio.create_task(self._kopiuj_tekst_dok())
                    ),
                    ft.Button(
                        content=ft.Text("TXT", size=12),
                        expand=True,
                        style=ft.ButtonStyle(padding=0),
                        on_click=lambda e: asyncio.create_task(self._eksportuj_dok("txt"))
                    ),
                    ft.Button(
                        content=ft.Text("Word", size=12),
                        expand=True,
                        style=ft.ButtonStyle(padding=0),
                        on_click=lambda e: asyncio.create_task(self._eksportuj_dok("docx"))
                    ),
                    ft.Button(
                        content=ft.Text("Excel", size=12),
                        expand=True,
                        style=ft.ButtonStyle(padding=0),
                        on_click=lambda e: asyncio.create_task(self._eksportuj_dok("xlsx"))
                    ),
                ], spacing=4, alignment=ft.MainAxisAlignment.SPACE_BETWEEN)
            ]
        )

    def _pobierz_deltas_dok(self, e):
        dx = getattr(e, "delta", None)
        if dx is not None:
            return e.delta.x, e.delta.y
        loc = getattr(e, "local_delta", None)
        if loc is not None:
            return e.local_delta.x, e.local_delta.y
        return getattr(e, "delta_x", 0.0), getattr(e, "delta_y", 0.0)

    def _aktualizuj_kadrowanie_dok(self):
        w = float(self.SZEROKOSC_DOK_DOC)
        h = float(self.WYSOKOSC_DOK_DOC)
        min_rozmiar = 40.0

        self.crop_doc_x1 = max(0.0, min(self.crop_doc_x1, w - min_rozmiar))
        self.crop_doc_y1 = max(0.0, min(self.crop_doc_y1, h - min_rozmiar))
        self.crop_doc_x2 = max(self.crop_doc_x1 + min_rozmiar, min(self.crop_doc_x2, w))
        self.crop_doc_y2 = max(self.crop_doc_y1 + min_rozmiar, min(self.crop_doc_y2, h))

        x1, y1, x2, y2 = self.crop_doc_x1, self.crop_doc_y1, self.crop_doc_x2, self.crop_doc_y2
        pol_uchwytu = self.UCHWYT_ROZMIAR_DOK_DOC / 2.0

        self.maska_doc_gora.width = w
        self.maska_doc_gora.height = y1

        self.maska_doc_dol.width = w
        self.maska_doc_dol.height = max(0.0, h - y2)

        self.maska_doc_lewo.top = y1
        self.maska_doc_lewo.height = max(0.0, y2 - y1)
        self.maska_doc_lewo.width = x1

        self.maska_doc_prawo.top = y1
        self.maska_doc_prawo.height = max(0.0, y2 - y1)
        self.maska_doc_prawo.width = max(0.0, w - x2)

        self.strefa_srodka_dok.top = y1
        self.strefa_srodka_dok.left = x1
        self.strefa_srodka_dok.width = max(0.0, x2 - x1)
        self.strefa_srodka_dok.height = max(0.0, y2 - y1)

        # Centrowanie uchwytów
        self.uchwyt_doc_lt.left = max(0.0, x1 - pol_uchwytu)
        self.uchwyt_doc_lt.top = max(0.0, y1 - pol_uchwytu)

        self.uchwyt_doc_rt.left = min(w - self.UCHWYT_ROZMIAR_DOK_DOC, x2 - pol_uchwytu)
        self.uchwyt_doc_rt.top = max(0.0, y1 - pol_uchwytu)

        self.uchwyt_doc_lb.left = max(0.0, x1 - pol_uchwytu)
        self.uchwyt_doc_lb.top = min(h - self.UCHWYT_ROZMIAR_DOK_DOC, y2 - pol_uchwytu)

        self.uchwyt_doc_rb.left = min(w - self.UCHWYT_ROZMIAR_DOK_DOC, x2 - pol_uchwytu)
        self.uchwyt_doc_rb.top = min(h - self.UCHWYT_ROZMIAR_DOK_DOC, y2 - pol_uchwytu)

        self.page.update()

    def _przesun_uchwyt_dok(self, ktory: str, dx: float, dy: float):
        self.kadr_doc_zmieniony = True
        if ktory == "lt":
            self.crop_doc_x1 += dx
            self.crop_doc_y1 += dy
        elif ktory == "rt":
            self.crop_doc_x2 += dx
            self.crop_doc_y1 += dy
        elif ktory == "lb":
            self.crop_doc_x1 += dx
            self.crop_doc_y2 += dy
        elif ktory == "rb":
            self.crop_doc_x2 += dx
            self.crop_doc_y2 += dy
        self._aktualizuj_kadrowanie_dok()

    def _przesun_caly_kadr_dok(self, dx: float, dy: float):
        self.kadr_doc_zmieniony = True
        szer_k = self.crop_doc_x2 - self.crop_doc_x1
        wys_k = self.crop_doc_y2 - self.crop_doc_y1

        self.crop_doc_x1 += dx
        self.crop_doc_y1 += dy
        self.crop_doc_x2 = self.crop_doc_x1 + szer_k
        self.crop_doc_y2 = self.crop_doc_y1 + wys_k

        if self.crop_doc_x1 < 0:
            self.crop_doc_x1 = 0.0
            self.crop_doc_x2 = szer_k
        if self.crop_doc_y1 < 0:
            self.crop_doc_y1 = 0.0
            self.crop_doc_y2 = wys_k
        if self.crop_doc_x2 > self.SZEROKOSC_DOK_DOC:
            self.crop_doc_x2 = float(self.SZEROKOSC_DOK_DOC)
            self.crop_doc_x1 = self.crop_doc_x2 - szer_k
        if self.crop_doc_y2 > self.WYSOKOSC_DOK_DOC:
            self.crop_doc_y2 = float(self.WYSOKOSC_DOK_DOC)
            self.crop_doc_y1 = self.crop_doc_y2 - wys_k

        self._aktualizuj_kadrowanie_dok()

    def _dopasuj_pola_robocze_pod_obraz_dok(self, sciezka: str):
        try:
            with Image.open(sciezka) as img:
                w_orig, h_orig = img.size
        except Exception:
            w_orig, h_orig = 1000, 1400

        self._ustaw_orientacje_sync([ft.DeviceOrientation.PORTRAIT_UP])

        max_w = self.MAX_SZEROKOSC_ROBOCZA_DOC
        max_h = self.MAX_WYSOKOSC_ROBOCZA_DOC

        proporcja = w_orig / max(1, h_orig)
        if (max_w / max_h) > proporcja:
            h_ramki = max_h
            w_ramki = round(max_h * proporcja)
        else:
            w_ramki = max_w
            h_ramki = round(max_w / proporcja)

        self.SZEROKOSC_DOK_DOC = max(140, int(w_ramki))
        self.WYSOKOSC_DOK_DOC = max(140, int(h_ramki))

        self.ramka_kadrowania_dok.width = self.SZEROKOSC_DOK_DOC
        self.ramka_kadrowania_dok.height = self.WYSOKOSC_DOK_DOC
        self.img_pelny_podglad_dok.width = self.SZEROKOSC_DOK_DOC
        self.img_pelny_podglad_dok.height = self.WYSOKOSC_DOK_DOC

        self.crop_doc_x1 = 0.0
        self.crop_doc_y1 = 0.0
        self.crop_doc_x2 = float(self.SZEROKOSC_DOK_DOC)
        self.crop_doc_y2 = float(self.WYSOKOSC_DOK_DOC)
        self.kadr_doc_zmieniony = False
        self._aktualizuj_kadrowanie_dok()

    def _ustaw_orientacje_sync(self, orientacje: list):
        plat = getattr(self.page, "platform", None)
        if plat in (ft.PagePlatform.ANDROID, ft.PagePlatform.IOS, "android", "ios"):
            try:
                res = self.page.set_allowed_device_orientations(orientacje)
                if asyncio.iscoroutine(res):
                    asyncio.create_task(res)
            except Exception:
                pass

    def _napraw_orientacje_exif(self, sciezka: str) -> str:
        try:
            with Image.open(sciezka) as img:
                img_poprawiony = ImageOps.exif_transpose(img)
                if img_poprawiony:
                    img_poprawiony.save(sciezka, quality=95)
        except Exception:
            pass
        return sciezka

    def ustaw_stan_foto_dok(self, czy_ma: bool):
        self.kontener_podgladu_dok.visible = czy_ma
        self.btn_otworz_kadr_dok.visible = czy_ma
        self.btn_start_dok.visible = czy_ma
        self.btn_usun_foto_dok.visible = czy_ma
        self.page.update()

    def ustaw_nowy_obraz_dok(self, sciezka: str):
        nowa_sciezka = os.path.join(config.KATALOG_DANYCH, f"img_dok_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg")
        try:
            shutil.copyfile(sciezka, nowa_sciezka)
        except Exception:
            nowa_sciezka = sciezka

        self._napraw_orientacje_exif(nowa_sciezka)

        self.zdjecie_dok["sciezka"] = nowa_sciezka
        self.historia_dok = [(nowa_sciezka, None)]
        self.aktywny_filtr_dok = None
        self._odswiez_styl_przyciskow_filtrow_dok()
        self.btn_doc_cofnij.disabled = True

        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = nowa_sciezka
        self.img_pelny_podglad_dok.src = nowa_sciezka

        self.ustaw_stan_foto_dok(True)
        self.status_dok.value = "Zdjęcie gotowe. Kliknij w podgląd, aby dopasować kadr."
        self.status_dok.color = ft.Colors.CYAN_ACCENT
        self.page.update()

    async def otworz_pelny_podglad_dok(self, e):
        if not self.zdjecie_dok["sciezka"] or not os.path.exists(self.zdjecie_dok["sciezka"]):
            return

        self._dopasuj_pola_robocze_pod_obraz_dok(self.zdjecie_dok["sciezka"])
        self.img_pelny_podglad_dok.src = self.zdjecie_dok["sciezka"]

        self.kolumna_glowna_dok.visible = False
        self.widok_kadrowania_dok.visible = True
        self.page.scroll = None
        self.page.update()

    async def _zamknij_pelny_ekran_dok(self, zapisz=False):
        self.widok_kadrowania_dok.visible = False
        self.kolumna_glowna_dok.visible = True
        self.page.scroll = ft.ScrollMode.AUTO

        self._ustaw_orientacje_sync([ft.DeviceOrientation.PORTRAIT_UP])
        self.page.update()

    async def obroc_dok(self, kat: int):
        if not self.zdjecie_dok["sciezka"] or not os.path.exists(self.zdjecie_dok["sciezka"]):
            return

        sciezka_akt = self.zdjecie_dok["sciezka"]
        self.historia_dok.append((sciezka_akt, self.aktywny_filtr_dok))
        self.btn_doc_cofnij.disabled = False

        loop = asyncio.get_running_loop()
        nowa_sciezka = await loop.run_in_executor(
            None, core.obroc_plik_graficzny, sciezka_akt, kat, "rot_dok"
        )

        self.zdjecie_dok["sciezka"] = str(nowa_sciezka)
        self.img_pelny_podglad_dok.src_base64 = None
        self.img_pelny_podglad_dok.src = str(nowa_sciezka)
        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = str(nowa_sciezka)

        self._dopasuj_pola_robocze_pod_obraz_dok(str(nowa_sciezka))
        self.page.update()

    async def _przelacz_filtr_pelny_dok(self, typ: str):
        if not self.zdjecie_dok["sciezka"] or not os.path.exists(self.zdjecie_dok["sciezka"]):
            return

        if self.aktywny_filtr_dok == typ:
            self._cofnij_krok_pelny_dok()
            return

        sciezka_akt = self.zdjecie_dok["sciezka"]
        self.historia_dok.append((sciezka_akt, self.aktywny_filtr_dok))
        self.btn_doc_cofnij.disabled = False

        loop = asyncio.get_running_loop()
        nowa_sciezka = await loop.run_in_executor(
            None, core.filtruj_plik_graficzny, sciezka_akt, typ
        )

        self.zdjecie_dok["sciezka"] = nowa_sciezka
        self.aktywny_filtr_dok = typ
        self.img_pelny_podglad_dok.src_base64 = None
        self.img_pelny_podglad_dok.src = nowa_sciezka
        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = nowa_sciezka

        self._odswiez_styl_przyciskow_filtrow_dok()
        self.page.update()

    def _odswiez_styl_przyciskow_filtrow_dok(self):
        aktywny_kolor = ft.Colors.BLUE_700
        zwykly_kolor = ft.Colors.GREY_800

        self.btn_doc_f_bw.style.bgcolor = aktywny_kolor if self.aktywny_filtr_dok == "bw" else zwykly_kolor
        self.btn_doc_f_szary.style.bgcolor = aktywny_kolor if self.aktywny_filtr_dok == "szary" else zwykly_kolor
        self.btn_doc_f_wyostrz.style.bgcolor = aktywny_kolor if self.aktywny_filtr_dok == "wyostrz" else zwykly_kolor

    async def _zatwierdz_krok_pelny_dok(self):
        if not self.zdjecie_dok["sciezka"] or not os.path.exists(self.zdjecie_dok["sciezka"]):
            return

        proc_lewo = (self.crop_doc_x1 / self.SZEROKOSC_DOK_DOC) * 100.0
        proc_gora = (self.crop_doc_y1 / self.WYSOKOSC_DOK_DOC) * 100.0
        proc_prawo = ((self.SZEROKOSC_DOK_DOC - self.crop_doc_x2) / self.SZEROKOSC_DOK_DOC) * 100.0
        proc_dol = ((self.WYSOKOSC_DOK_DOC - self.crop_doc_y2) / self.WYSOKOSC_DOK_DOC) * 100.0

        if proc_lewo <= 0.5 and proc_gora <= 0.5 and proc_prawo <= 0.5 and proc_dol <= 0.5:
            return

        sciezka_akt = self.zdjecie_dok["sciezka"]
        self.historia_dok.append((sciezka_akt, self.aktywny_filtr_dok))
        self.btn_doc_cofnij.disabled = False

        loop = asyncio.get_running_loop()
        wyciety_plik = await loop.run_in_executor(
            None, core.kadruj_plik_graficzny, sciezka_akt,
            proc_lewo, proc_gora, proc_prawo, proc_dol
        )

        self.zdjecie_dok["sciezka"] = wyciety_plik
        self.img_pelny_podglad_dok.src_base64 = None
        self.img_pelny_podglad_dok.src = wyciety_plik
        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = wyciety_plik

        self._dopasuj_pola_robocze_pod_obraz_dok(wyciety_plik)

    def _cofnij_krok_pelny_dok(self):
        if not self.historia_dok:
            return
        poprzednia_sciezka, poprzedni_filtr = self.historia_dok.pop()
        self.zdjecie_dok["sciezka"] = poprzednia_sciezka
        self.aktywny_filtr_dok = poprzedni_filtr

        self.img_pelny_podglad_dok.src_base64 = None
        self.img_pelny_podglad_dok.src = poprzednia_sciezka
        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = poprzednia_sciezka

        self._dopasuj_pola_robocze_pod_obraz_dok(poprzednia_sciezka)
        self._odswiez_styl_przyciskow_filtrow_dok()
        self.btn_doc_cofnij.disabled = (len(self.historia_dok) == 0)
        self.page.update()

    async def _wyslij_pelny_ekran_dok(self):
        if self.kadr_doc_zmieniony:
            await self._zatwierdz_krok_pelny_dok()
        await self._zamknij_pelny_ekran_dok(zapisz=True)
        await self.analizuj_dokument_dok(None)

    def usun_wybrane_zdjecie_dok(self):
        sciezka_pliku = self.zdjecie_dok.get("sciezka")
        if sciezka_pliku and os.path.exists(sciezka_pliku):
            try:
                os.remove(sciezka_pliku)
                self.ui.dopisz_log(f"Usunięto plik: {os.path.basename(sciezka_pliku)}")
            except Exception as err:
                self.ui.dopisz_log(f"Błąd usuwania pliku: {err}", ft.Colors.AMBER)

        self.zdjecie_dok["sciezka"] = None
        self.podglad_dok.src_base64 = None
        self.podglad_dok.src = config.PUSTY_OBRAZ
        self.ustaw_stan_foto_dok(False)
        self.status_dok.value = "Wybierz zdjęcie dokumentu biurowego lub zrób nowe."
        self.status_dok.color = ft.Colors.GREY_300
        self.page.update()

    async def otworz_galerie_dok(self):
        try:
            pliki = await self.pickery["foto"].pick_files(allow_multiple=False, file_type=ft.FilePickerFileType.IMAGE)
            if pliki and len(pliki) > 0 and pliki[0].path:
                self.ustaw_nowy_obraz_dok(pliki[0].path)
        except Exception as e_pick:
            self.status_dok.value = f"Błąd wyboru pliku: {e_pick}"
            self.page.update()

    async def analizuj_dokument_dok(self, e):
        if not self.zdjecie_dok["sciezka"]:
            return

        self.btn_start_dok.disabled = True
        self.pasek_dok.visible = True
        self.status_dok.value = "Przygotowywanie dokumentu do odczytu..."
        self.status_dok.color = ft.Colors.ORANGE_ACCENT

        import time
        start_calkowity = time.perf_counter()
        self.ui.dopisz_log("=== ROZPOCZĘTO ODCZYT DOKUMENTU (1:1) ===", ft.Colors.CYAN)
        self.page.update()

        loop = asyncio.get_running_loop()
        sciezka_do_analizy = self.zdjecie_dok["sciezka"]
        waga_oryg_kb = round(os.path.getsize(sciezka_do_analizy) / 1024, 1)

        if self.kadr_doc_zmieniony:
            proc_lewo = (self.crop_doc_x1 / self.SZEROKOSC_DOK_DOC) * 100.0
            proc_gora = (self.crop_doc_y1 / self.WYSOKOSC_DOK_DOC) * 100.0
            proc_prawo = ((self.SZEROKOSC_DOK_DOC - self.crop_doc_x2) / self.SZEROKOSC_DOK_DOC) * 100.0
            proc_dol = ((self.WYSOKOSC_DOK_DOC - self.crop_doc_y2) / self.WYSOKOSC_DOK_DOC) * 100.0
            sciezka_do_analizy = await loop.run_in_executor(
                None, core.kadruj_plik_graficzny, self.zdjecie_dok["sciezka"],
                proc_lewo, proc_gora, proc_prawo, proc_dol
            )
            self.ui.dopisz_log(f"✂️ Zastosowano kadr: L={proc_lewo:.0f}%, T={proc_gora:.0f}%, R={proc_prawo:.0f}%, B={proc_dol:.0f}%", ft.Colors.GREY_400)

        prompt_dok = (
            "Jesteś precyzyjnym systemem OCR do dokumentów biurowych, faktur i specyfikacji.\n"
            "Przepisz CAŁY tekst widoczny na obrazie z zachowaniem układu:\n"
            "1. Jeżeli na obrazie występuje tabela, zestawienie kolumnowe lub lista z cenami/ilościami, "
            "zapisz ją BEZWZGLĘDNIE jako standardową tabelę Markdown (używając pionowych kresek '|' oraz nagłówków '|---|').\n"
            "2. Zachowaj oryginalne akapity, tytuły i odstępy.\n"
            "3. Nie dodawaj żadnych własnych komentarzy, wstępów ani znaczników ```markdown. Zwróć sam czysty tekst z tabelami."
        )

        try:
            cfg = config.wczytaj_konfiguracje()
            uzywa_chmury = cfg.get("use_cloud", True)
            wymiar = cfg.get("image_resolution", 1800)

            if not uzywa_chmury:
                ip_lokalne = cfg.get("local_ip", "192.168.1.154").strip()
                port_str = cfg.get("local_port", "1234").strip()
                mac_adres = cfg.get("wol_mac", "").strip()
                port_lokalny = int(port_str) if port_str.isdigit() else 1234

                self.ui.dopisz_log(f"Sprawdzanie stanu serwera LM Studio ({ip_lokalne}:{port_lokalny})...")
                serwer_zyje = await core.sprawdz_port_tcp(ip_lokalne, port_lokalny, timeout=3.0)
                if not serwer_zyje:
                    self.ui.dopisz_log("Serwer lokalny nie odpowiada. Wysyłanie WoL...", ft.Colors.AMBER)
                    try:
                        await loop.run_in_executor(None, core.wyslij_wol, mac_adres, ip_lokalne)
                    except Exception:
                        pass
                    await asyncio.sleep(15)

            try:
                with Image.open(sciezka_do_analizy) as img:
                    w_px, h_px = img.size
            except Exception:
                w_px, h_px = "?", "?"

            base64_image = await loop.run_in_executor(None, core.kompresuj_do_base64, sciezka_do_analizy, wymiar)
            waga_b64_kb = round(len(base64_image) * 0.75 / 1024, 1)
            self.ui.dopisz_log(f"📷 Kadr wejściowy: {w_px}x{h_px}px | Waga: {waga_b64_kb} KB (oryginał: {waga_oryg_kb} KB)", ft.Colors.GREY_400)

            if uzywa_chmury:
                model_nazwa = cfg.get("gemini_model", "gemini-2.5-flash")
                url = "https://generativelanguage.googleapis.com/v1beta/openai/chat/completions"
                headers = {"Authorization": f"Bearer {cfg.get('gemini_api_key', '')}", "Content-Type": "application/json"}
                payload = {
                    "model": model_nazwa,
                    "messages": [{"role": "user", "content": [{"type": "text", "text": prompt_dok}, {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}]}],
                    "temperature": 0.0,
                    "max_tokens": 8192
                }
                cel_logu = f"Google Cloud ({model_nazwa})"
            else:
                model_nazwa = cfg.get("local_model", "qwen3-vl-8b-instruct").strip()
                url = f"http://{cfg.get('local_ip')}:{cfg.get('local_port')}/v1/chat/completions"
                headers = {"Content-Type": "application/json"}
                if cfg.get("local_api_key"):
                    headers["Authorization"] = f"Bearer {cfg.get('local_api_key')}"
                payload = {
                    "model": model_nazwa,
                    "messages": [{"role": "user", "content": [{"type": "text", "text": prompt_dok}, {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}]}],
                    "temperature": 0.0,
                    "chat_template_kwargs": {"enable_thinking": False}
                }
                cel_logu = f"LM Studio ({model_nazwa})"

            self.status_dok.value = "Odczytywanie dokumentu przez AI..."
            self.ui.dopisz_log(f"🌐 Wysyłanie żądania do: {cel_logu}...")
            self.page.update()

            max_prob = 4
            opoznienie_poczatkowe = 2.0
            odpowiedz = None
            timeout_cfg = httpx.Timeout(10.0, read=300.0)

            start_siec = time.perf_counter()
            async with httpx.AsyncClient(timeout=timeout_cfg, verify=True) as client:
                for proba in range(max_prob):
                    if proba > 0:
                        self.ui.dopisz_log(f"Ponawianie zapytania (próba {proba + 1}/{max_prob})...", ft.Colors.AMBER)
                    odpowiedz = await client.post(url, headers=headers, json=payload)

                    if odpowiedz.status_code in [503, 429]:
                        if proba < max_prob - 1:
                            czas_oczekiwania = opoznienie_poczatkowe * (2 ** proba)
                            self.status_dok.value = f"Serwer zajęty ({odpowiedz.status_code}). Ponawianie za {czas_oczekiwania:.1f}s..."
                            self.status_dok.color = ft.Colors.AMBER_ACCENT
                            self.page.update()
                            await asyncio.sleep(czas_oczekiwania)
                            continue

                    odpowiedz.raise_for_status()
                    break

            czas_siec = time.perf_counter() - start_siec
            dane_odp = odpowiedz.json()
            wybor = dane_odp["choices"][0]
            surowy_tekst = wybor["message"]["content"].strip()
            powod_konca = wybor.get("finish_reason")

            uzycie = dane_odp.get("usage") or {}
            in_tok = uzycie.get("prompt_tokens", 0)
            out_tok = uzycie.get("completion_tokens", 0)
            think_tok = (uzycie.get("completion_tokens_details") or {}).get("reasoning_tokens")

            predkosc = (out_tok / czas_siec) if (czas_siec > 0 and out_tok > 0) else 0.0

            log_wydajnosc = f"⏱️ Czas AI: {czas_siec:.2f}s | Tokeny: {in_tok} wejście / {out_tok} wyjście"
            if think_tok:
                log_wydajnosc += f" (myślenie: {think_tok})"
            if predkosc > 0:
                log_wydajnosc += f" | Prędkość: {predkosc:.1f} tok/s"
            self.ui.dopisz_log(log_wydajnosc, ft.Colors.CYAN)

            if powod_konca == "length":
                self.ui.dopisz_log("⚠️ OSTRZEŻENIE: Odpowiedź ucięta po limicie tokenów!", ft.Colors.RED)

            linie = surowy_tekst.splitlines()
            wiersze_tabeli = sum(1 for l in linie if l.strip().startswith('|') and l.strip().endswith('|'))
            slowa = len(surowy_tekst.split())
            znaki = len(surowy_tekst)

            self.ui.dopisz_log(
                f"📝 Odczytano: {znaki} znaków | {slowa} słów | {len(linie)} linii"
                + (f" | 📊 Wiersze tabeli: {wiersze_tabeli}" if wiersze_tabeli > 0 else ""),
                ft.Colors.GREEN
            )

            czas_calkowity = time.perf_counter() - start_calkowity
            self.ui.dopisz_log(f"✅ Sukces w czasie {czas_calkowity:.2f}s.", ft.Colors.GREEN)

            self.txt_edytor_dok.value = surowy_tekst
            self.ostatni_wynik_dok["tekst"] = surowy_tekst
            self.btn_wroc_wynik_dok.visible = True

            self.status_dok.value = "✅ Dokument został pomyślnie odczytany."
            self.status_dok.color = ft.Colors.GREEN_ACCENT
            self.ui.bezpiecznie_otworz_dialog(self.dlg_wynik_dok)

        except httpx.HTTPStatusError as http_err:
            status = http_err.response.status_code
            tresc = http_err.response.text[:350]
            self.ui.dopisz_log(f"Błąd HTTP {status}: {tresc}", ft.Colors.RED)
            if status in (401, 403) or "API_KEY_INVALID" in tresc:
                self.status_dok.value = "Błąd autoryzacji API"
                self.ui.pokaz_okno_bledu("🔑 Błąd autoryzacji", "Klucz API jest nieprawidłowy lub brak uprawnień. Sprawdź ustawienia.")
            elif status == 429:
                self.status_dok.value = "Limit zapytań wyczerpany"
                self.ui.pokaz_okno_bledu("⏳ Limit zapytań wyczerpany", "Zbyt wiele zapytań w krótkim czasie. Odczekaj 30 sekund.")
            elif status >= 500:
                self.status_dok.value = "Serwer AI niedostępny"
                self.ui.pokaz_okno_bledu("⏳ Serwer AI niedostępny", f"Serwer zwrócił kod {status}. Odczekaj chwilę i spróbuj ponownie.")
            else:
                self.status_dok.value = f"Błąd zapytania HTTP {status}"
                self.ui.pokaz_okno_bledu("❌ Błąd zapytania HTTP", f"Status: {status}\n{tresc}")
        except httpx.TimeoutException:
            self.status_dok.value = "Limit czasu przekroczony"
            self.ui.pokaz_okno_bledu("⏳ Limit czasu", "Model nie odpowiedział w wyznaczonym czasie.")
        except Exception as err:
            self.ui.dopisz_log(f"Błąd odczytu: {err}", ft.Colors.RED)
            self.status_dok.value = f"Błąd: {err}"
            self.status_dok.color = ft.Colors.RED_ACCENT
            self.ui.pokaz_okno_bledu("Błąd odczytu dokumentu", str(err))
        finally:
            self.pasek_dok.visible = False
            self.btn_start_dok.disabled = False
            self.page.update()

    async def _kopiuj_tekst_dok(self):
        tekst = self.txt_edytor_dok.value or ""
        await self.page.set_clipboard(tekst)
        self.ui.dopisz_log("Skopiowano tekst dokumentu do schowka systemowego.", ft.Colors.CYAN)

    async def _eksportuj_dok(self, format_pliku: str):
        tekst = self.txt_edytor_dok.value or ""
        if not tekst.strip():
            self.ui.pokaz_okno_bledu("Brak tekstu", "Brak treści do wyeksportowania.")
            return

        stem = f"dok_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        sciezka_wyjsciowa = os.path.join(config.KATALOG_DANYCH, f"{stem}.{format_pliku}")

        try:
            if format_pliku == "txt":
                with open(sciezka_wyjsciowa, "w", encoding="utf-8") as f:
                    f.write(tekst)

            elif format_pliku == "docx":
                import docx
                doc = docx.Document()
                for linia in tekst.splitlines():
                    doc.add_paragraph(linia)
                doc.save(sciezka_wyjsciowa)

            elif format_pliku == "xlsx":
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Dokument"
                for r_idx, linia in enumerate(tekst.splitlines(), start=1):
                    if "|" in linia:
                        czesci = [c.strip() for c in linia.split("|") if c.strip()]
                        for c_idx, val in enumerate(czesci, start=1):
                            ws.cell(row=r_idx, column=c_idx, value=val)
                    else:
                        ws.cell(row=r_idx, column=1, value=linia)
                wb.save(sciezka_wyjsciowa)

            self.ui.dopisz_log(f"Zapisano plik: {os.path.basename(sciezka_wyjsciowa)}", ft.Colors.GREEN)
            await self._bezpiecznie_udostepnij_plik(sciezka_wyjsciowa, f"Dokument {format_pliku.upper()}")
        except Exception as err:
            self.ui.dopisz_log(f"Błąd eksportu {format_pliku}: {err}", ft.Colors.RED)
            self.ui.pokaz_okno_bledu("Błąd eksportu", str(err))

    async def _bezpiecznie_udostepnij_plik(self, sciezka: str, tytul: str):
        try:
            if hasattr(self.ui.serwis_udostepniania, "share_files"):
                try:
                    await self.ui.serwis_udostepniania.share_files(
                        [ft.ShareFile.from_path(sciezka)],
                        text=tytul
                    )
                except Exception:
                    await self.ui.serwis_udostepniania.share_files([sciezka])
        except Exception as e_share:
            self.ui.dopisz_log(f"Nie udało się otworzyć systemowego okna udostępniania: {e_share}", ft.Colors.AMBER)
